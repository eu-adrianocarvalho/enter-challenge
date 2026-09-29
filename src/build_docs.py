"""Gera docs/index.html: o site que resume o projeto para a reunião, no visual da Enter (getenter.ai).
O conteúdo é dividido em quatro tópicos (Apresentação, O trabalho, Uso, Arquivos); a barra lateral mostra
os tópicos e expande só os itens do tópico aberto. As seções vêm de docs/0*.md e do relatório, os
números da última execução são lidos de Output/ e o catálogo traz prévia e download de cada arquivo.
Rode depois de gerar a carta: python src/build_docs.py."""
from __future__ import annotations

import csv
import html
import json
import re
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import quote

import markdown
import pypdfium2 as pdfium
from docx import Document

from docsite.rivet_view import RIVET_CSS, RIVET_JS, rivet_section
from xp_letter.config import REPO_ROOT

DOCS = REPO_ROOT / "docs"
ASSETS = DOCS / "assets"
OUTPUT = REPO_ROOT / "Output"
TEXT_PREVIEW_LIMIT = 60_000
CSV_PREVIEW_ROWS = 60
AUTHOR = "Adriano da Silva de Carvalho"
REPO_URL = "https://github.com/eu-adrianocarvalho/enter-challenge"
MERMAID_URL = "https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.esm.min.mjs"
MERMAID_BLOCK = re.compile(r'<pre><code class="language-mermaid">(.*?)</code></pre>', flags=re.DOTALL)
FONTS = ("https://fonts.googleapis.com/css2?family=Geist:wght@400;500;600&family=Geist+Mono:wght@400;500"
         "&family=Noto+Serif:wght@300;400&display=swap")


@dataclass(frozen=True)
class Topic:
    id: str
    label: str
    summary: str


@dataclass(frozen=True)
class Section:
    source: str
    anchor: str
    nav: str
    topic: str


TOPICS = (
    Topic("apresentacao", "Apresentação", "O que foi entregue, os números e a carta final ao lado da v1."),
    Topic("trabalho", "O trabalho", "O diagnóstico da v1, as melhorias, a arquitetura, os dados externos e as travas."),
    Topic("uso", "Uso", "Como rodar, o roteiro da reunião e o relatório de duas páginas."),
    Topic("arquivos", "Arquivos", "Prévia e download de tudo o que entrou e saiu do sistema."),
)

SECTIONS = (
    Section("01_visao_geral.md", "visao-geral", "Visão geral", "apresentacao"),
    Section("", "carta", "A carta e o brief", "apresentacao"),
    Section("02_diagnostico_v1.md", "diagnostico", "Diagnóstico da v1", "trabalho"),
    Section("03_melhorias_implementadas.md", "melhorias", "Melhorias implementadas", "trabalho"),
    Section("04_arquitetura_e_codigo.md", "arquitetura", "Arquitetura e código", "trabalho"),
    Section("", "rivet", "Grafos do Rivet", "trabalho"),
    Section("05_dados_externos.md", "dados-externos", "Dados externos", "trabalho"),
    Section("06_qualidade_e_travas.md", "qualidade", "Qualidade e travas", "trabalho"),
    Section("07_como_usar.md", "como-usar", "Como usar", "uso"),
    Section("08_roteiro_da_reuniao.md", "roteiro", "Roteiro da reunião", "uso"),
    Section("relatorio.md", "relatorio", "Relatório (2 páginas)", "uso"),
)


def code_files() -> tuple[tuple[str, str], ...]:
    sources = sorted(p for p in (REPO_ROOT / "src").rglob("*.py") if "__pycache__" not in p.parts)
    return tuple(
        (p.relative_to(REPO_ROOT).as_posix(), p.read_text(encoding="utf-8").split("\n", 1)[0].strip('"'))
        for p in sources
    )


def _listing(folder: str, pattern: str, description: str) -> tuple[tuple[str, str], ...]:
    return tuple((f"{folder}/{p.name}", description) for p in sorted((REPO_ROOT / folder).glob(pattern)))


FILE_GROUPS = (
    ("arquivos-entradas", "Entradas do desafio", (
        ("Input/XP - Albert_s portfolio.pdf", "Extrato da carteira do Albert (PDF original)"),
        ("Input/XP - Albert_s risk profile.pdf", "Perfil de risco do Albert (PDF original)"),
        ("Input/XP - Macro analysis.pdf", "Relatório Brasil Macro Mensal da XP, 06/02/2025"),
        ("Input/profitability_calc_wip.csv", "Preços atual e do mês anterior das ações (valores corretos)"),
        ("Input/profitability_calc_wip.xlsx", "Planilha de rentabilidade inacabada, com números corrompidos"),
        ("Input/XP - Albert_s portfolio.txt", "Texto do extrato entregue no desafio (colunas embaralhadas)"),
        ("Input/XP - Albert_s risk profile.txt", "Texto do perfil de risco"),
        ("Input/XP - Macro analysis.txt", "Texto do relatório macro entregue no desafio"),
    )),
    ("arquivos-saidas", "Saídas geradas (v2)", (
        ("Output/carta_albert_2025-05-07.pdf", "Carta para o cliente, 2 páginas"),
        ("Output/carta_albert_2025-05-07.docx", "Mesma carta, editável pelo assessor"),
        ("Output/brief_assessor_albert_2025-05-07.md", "Brief do assessor: status, travas, alertas, recomendações e custo"),
        ("Output/grafico_albert_2025-05-07.png", "Gráfico da carta"),
        ("Output/facts_albert_2025-05-07.json", "FACTS usados pela carta e o texto final"),
        ("Output/run_log_albert_2025-05-07.json", "Tokens, custo e tempo de cada chamada ao LLM"),
        ("docs/relatorio.pdf", "Relatório curto do desafio (2 páginas)"),
    )),
    ("arquivos-v1", "Primeira versão (v1)", (
        ("enter_challenge.rivet-project", "Grafo original do Rivet, intacto"),
        ("Output/output_letter.docx", "Carta gerada pela v1, intacta"),
    )),
    ("arquivos-rivet", "Workflow Rivet (v2)", (
        ("rivet/xp_monthly_letter.rivet-project", "Projeto com os 6 grafos (gerado por src/build_rivet.py)"),
        *_listing("rivet/prompts", "*.md", "Prompt"),
        *_listing("rivet/schemas", "*.json", "JSON schema da resposta"),
    )),
    ("arquivos-config", "Configuração", (
        ("config/settings.yaml", "Entradas, período, modelos, preços por token e limites"),
        ("config/allocation_moderate.yaml", "Faixas de alocação do perfil moderado (ilustrativas)"),
        ("config/research_shelf.yaml", "Produtos e ações que o motor pode recomendar (ilustrativos)"),
        ("config/fund_registry.yaml", "Fundos do extrato → CNPJ na CVM, conferidos à mão"),
    )),
    ("arquivos-dados", "Dados de mercado e LLM", (
        ("data/README.md", "O que é cada pasta de data/"),
        *_listing("data/market", "*", "Dado de mercado do período"),
        *_listing("data/llm", "*.json", "Resposta guardada de um grafo"),
    )),
    ("arquivos-evidencias", "Evidências", _listing("data/evidence", "*.json", "Resposta real do LLM que motivou uma trava")),
    ("arquivos-codigo", "Código", code_files() + (
        ("README.md", "README do repositório (em inglês)"),
        ("requirements.txt", "Dependências Python"),
        ("package.json", "Dependência Node (rivet-cli)"),
    )),
)

CSS = """
:root { --bg:#ffffff; --surface:#f3f3f3; --text:#000000; --text-2:rgba(0,0,0,.64); --stroke:rgba(0,0,0,.1);
  --stroke-strong:rgba(0,0,0,.2); --hover:rgba(0,0,0,.05); --inv-bg:#171717; --inv-text:#ffffff;
  --inv-text-2:rgba(255,255,255,.64); --inv-stroke:rgba(255,255,255,.16); --accent:#ffae35; --on-accent:#000000;
  --page-bg:#e8e8e8; --serif:"Noto Serif", Georgia, serif; --sans:"Geist", "Segoe UI", system-ui, sans-serif;
  --mono:"Geist Mono", Consolas, monospace; }
@media (prefers-color-scheme: dark) { :root:not([data-theme="light"]) { --bg:#0e0e0e; --surface:#171717; --text:#ffffff;
  --text-2:rgba(255,255,255,.64); --stroke:rgba(255,255,255,.1); --stroke-strong:rgba(255,255,255,.2);
  --hover:rgba(255,255,255,.06); --inv-bg:#1c1c1c; --inv-stroke:rgba(255,255,255,.12); --page-bg:#262626; } }
:root[data-theme="dark"] { --bg:#0e0e0e; --surface:#171717; --text:#ffffff; --text-2:rgba(255,255,255,.64);
  --stroke:rgba(255,255,255,.1); --stroke-strong:rgba(255,255,255,.2); --hover:rgba(255,255,255,.06); --inv-bg:#1c1c1c;
  --inv-stroke:rgba(255,255,255,.12); --page-bg:#262626; }
* { box-sizing:border-box; }
body { margin:0; background:var(--bg); color:var(--text); font:16px/1.6 var(--sans); letter-spacing:-.1px; }
.topbar { background:#000; color:#fff; font-size:13px; text-align:center; padding:11px 16px; }
.topbar a { color:var(--accent); text-transform:uppercase; letter-spacing:.2px; text-decoration:none; margin-left:14px; font-size:12px; }
.layout { display:grid; grid-template-columns:300px minmax(0,1fr); }
nav.sidebar { position:sticky; top:0; align-self:start; height:100vh; overflow-y:auto; padding:28px 22px;
  border-right:1px solid var(--stroke); background:var(--bg); display:flex; flex-direction:column; }
.brand { padding-bottom:22px; border-bottom:1px solid var(--stroke); margin-bottom:14px; }
.brand .mark { font-size:22px; letter-spacing:.5px; font-weight:500; }
.brand .mark b { display:inline-block; width:12px; height:12px; background:var(--text); margin-left:4px; vertical-align:1px; }
.brand .author { margin-top:14px; font-size:13px; line-height:1.4; }
.brand .author strong { display:block; font-weight:600; }
.brand .author a { font:11.5px var(--mono); color:var(--text-2); word-break:break-all; text-decoration:none; }
.brand .author a:hover { color:var(--text); text-decoration:underline; text-decoration-color:var(--accent); }
pre.mermaid { background:transparent; border:1px solid var(--stroke); border-radius:12px; padding:22px; text-align:center;
  max-height:none; overflow:auto; font-family:var(--sans); color:var(--text-2); margin:18px 0 22px; }
.brand small { display:block; color:var(--text-2); font-size:12px; letter-spacing:.2px; text-transform:uppercase; margin-top:6px; }
.nav-topic { border-bottom:1px solid var(--stroke); }
.topic-link { display:flex; gap:12px; align-items:center; padding:14px 6px; color:var(--text); text-decoration:none;
  text-transform:uppercase; letter-spacing:.2px; font-size:13px; }
.topic-link .num { font-family:var(--mono); color:var(--text-2); font-size:12px; }
.topic-link::after { content:"+"; margin-left:auto; color:var(--text-2); }
.nav-topic.open .topic-link::after { content:"–"; }
.nav-topic.open .topic-link { font-weight:600; }
.nav-topic.open .topic-link .num { color:var(--text); background:var(--accent); padding:1px 5px; border-radius:4px; }
.subnav { display:none; padding:0 0 12px 34px; }
.nav-topic.open .subnav { display:block; }
.subnav a { display:block; padding:5px 10px; margin:1px 0; border-radius:8px; color:var(--text-2); text-decoration:none;
  font-size:14px; border-left:2px solid transparent; }
.subnav a:hover { background:var(--hover); color:var(--text); }
.subnav a.active { color:var(--text); border-left-color:var(--accent); background:var(--hover); }
.theme-toggle { margin-top:24px; padding:12px 14px; border:1px solid var(--stroke-strong); border-radius:8px; background:transparent;
  color:var(--text); cursor:pointer; text-transform:uppercase; letter-spacing:.2px; font:12px var(--sans); }
main { padding:48px clamp(24px, 4vw, 72px) 96px; max-width:1560px; width:100%; }
.page[hidden] { display:none; }
.page-head { display:flex; align-items:baseline; gap:16px; border-bottom:1px solid var(--stroke); padding-bottom:18px; margin-bottom:8px; }
.page-head .num { font-family:var(--mono); font-size:14px; color:var(--text-2); }
.page-head h1 { font:300 clamp(34px, 3.4vw, 52px)/1.15 var(--serif); letter-spacing:-.2px; margin:0; }
.page-head p { margin:0 0 0 auto; color:var(--text-2); max-width:420px; font-size:14px; text-align:right; }
.hero { background:var(--inv-bg); color:var(--inv-text); border-radius:20px; padding:40px 44px; margin:28px 0 12px; }
.hero .eyebrow { text-transform:uppercase; letter-spacing:.2px; font-size:12px; color:var(--inv-text-2); }
.hero h2 { font:300 clamp(30px, 3vw, 46px)/1.15 var(--serif); margin:10px 0 12px; border:none; padding:0; letter-spacing:-.2px; }
.hero h2 em { font-style:normal; color:var(--accent); }
.hero p { color:var(--inv-text-2); max-width:760px; margin:0; }
.kpis { display:grid; grid-template-columns:repeat(4, minmax(0,1fr)); margin-top:32px; border-top:1px solid var(--inv-stroke); }
.kpi { padding:20px 18px 18px 0; border-bottom:1px solid var(--inv-stroke); }
.kpi + .kpi { padding-left:18px; border-left:1px solid var(--inv-stroke); }
.kpi:nth-child(4n+1) { padding-left:0; border-left:none; }
.kpi span { display:block; font-size:12px; text-transform:uppercase; letter-spacing:.2px; color:var(--inv-text-2); }
.kpi strong { display:block; font:300 30px/1.2 var(--serif); margin-top:6px; white-space:nowrap; }
.hero .actions { display:flex; gap:12px; margin-top:28px; flex-wrap:wrap; }
.btn { display:inline-flex; align-items:center; gap:10px; background:var(--accent); color:var(--on-accent); border-radius:8px;
  padding:14px 20px; text-transform:uppercase; letter-spacing:.2px; font-size:13px; text-decoration:none; border:none; cursor:pointer; }
.btn.secondary { background:transparent; color:inherit; border:1px solid var(--stroke-strong); }
.hero .btn.secondary { border-color:var(--inv-stroke); color:var(--inv-text); }
section.doc { padding:40px 0 8px; border-bottom:1px solid var(--stroke); scroll-margin-top:24px; }
section.doc > h1, section.doc .doc-title { font:300 clamp(28px, 2.4vw, 38px)/1.2 var(--serif); letter-spacing:-.1px; margin:0 0 18px; }
section.doc h2 { font:300 25px/1.25 var(--serif); margin:34px 0 12px; }
section.doc h3 { font:600 17px/1.4 var(--sans); margin:24px 0 8px; }
p, li { max-width:980px; }
table { border-collapse:collapse; width:100%; margin:14px 0 20px; font-size:14.5px; display:block; overflow-x:auto; }
th { text-align:left; font-weight:500; font-size:12px; text-transform:uppercase; letter-spacing:.2px; color:var(--text-2);
  padding:10px 12px; border-bottom:1px solid var(--stroke-strong); white-space:nowrap; }
td { padding:10px 12px; border-bottom:1px solid var(--stroke); vertical-align:top; }
tbody tr:hover td { background:var(--hover); }
code { font-family:var(--mono); background:var(--surface); padding:1px 6px; border-radius:4px; font-size:13px; }
pre { font-family:var(--mono); background:var(--surface); padding:16px 18px; border-radius:12px; overflow:auto; font-size:12.5px;
  line-height:1.5; max-height:560px; border:1px solid var(--stroke); }
pre code { background:none; padding:0; }
blockquote { margin:14px 0; padding:14px 20px; border-left:3px solid var(--accent); background:var(--surface); border-radius:0 12px 12px 0; }
a { color:inherit; text-underline-offset:3px; }
a:hover { color:var(--text); text-decoration-color:var(--accent); }
strong { font-weight:600; }
.two-col { display:grid; grid-template-columns:minmax(0,1.15fr) minmax(0,.85fr); gap:28px; }
.col-title { display:flex; align-items:center; justify-content:space-between; margin:0 0 10px; }
.col-title span { text-transform:uppercase; letter-spacing:.2px; font-size:12px; color:var(--text-2); }
.col-title a { font-size:12px; text-transform:uppercase; letter-spacing:.2px; }
.pages { max-height:880px; overflow:auto; border-radius:12px; background:var(--page-bg); padding:14px; border:1px solid var(--stroke); }
.pages img.page { display:block; width:100%; margin:0 0 14px; border-radius:4px; box-shadow:0 1px 4px rgba(0,0,0,.18); }
.v1-letter { font-size:14.5px; color:var(--text-2); max-height:880px; overflow:auto; border:1px solid var(--stroke);
  border-radius:12px; padding:18px 22px; background:var(--surface); }
.frame { width:100%; height:780px; border:1px solid var(--stroke); border-radius:12px; background:#fff; }
.brief { border:1px solid var(--stroke); border-radius:12px; padding:6px 26px 18px; margin-top:12px; }
.file-tools { margin:6px 0 18px; }
.file-tools input { width:100%; padding:14px 16px; border:1px solid var(--stroke-strong); border-radius:8px; background:var(--bg);
  color:var(--text); font:15px var(--sans); }
section.files h2 { font:300 25px/1.25 var(--serif); margin:28px 0 8px; }
details.file { border-bottom:1px solid var(--stroke); }
details.file > summary { list-style:none; cursor:pointer; padding:13px 6px; display:grid;
  grid-template-columns:18px minmax(220px, 1.1fr) minmax(0, 1.3fr) 70px auto; gap:16px; align-items:center; }
details.file > summary:hover { background:var(--hover); }
details.file > summary::-webkit-details-marker { display:none; }
details.file > summary::before { content:"+"; color:var(--text-2); font-family:var(--mono); }
details.file[open] > summary::before { content:"–"; }
.fname { font-family:var(--mono); font-size:13px; word-break:break-all; }
.fdesc { color:var(--text-2); font-size:14px; }
.fsize { color:var(--text-2); font-size:12px; font-family:var(--mono); text-align:right; }
.dl { font-size:11px; text-transform:uppercase; letter-spacing:.2px; text-decoration:none; padding:7px 12px; border-radius:8px;
  background:var(--accent); color:var(--on-accent); }
.preview { padding:4px 6px 22px 40px; }
.preview img { max-width:100%; border:1px solid var(--stroke); border-radius:8px; background:#fff; }
.note { color:var(--text-2); font-size:14px; font-style:italic; }
.pager { display:flex; justify-content:space-between; gap:16px; margin-top:48px; }
.pager a { flex:1; max-width:48%; border:1px solid var(--stroke-strong); border-radius:12px; padding:18px 22px; text-decoration:none; }
.pager a:hover { background:var(--hover); }
.pager a.next { margin-left:auto; text-align:right; }
.pager small { display:block; text-transform:uppercase; letter-spacing:.2px; font-size:12px; color:var(--text-2); }
.pager strong { font:300 22px/1.3 var(--serif); }
@media (max-width: 1100px) { .two-col { grid-template-columns:1fr; } .kpis { grid-template-columns:repeat(2, minmax(0,1fr)); }
  .kpi:nth-child(2n+1) { padding-left:0; border-left:none; } .kpi:nth-child(4n+1) { padding-left:0; } }
@media (max-width: 860px) {
  .layout { grid-template-columns:1fr; }
  nav.sidebar { position:static; height:auto; border-right:none; border-bottom:1px solid var(--stroke); }
  .theme-toggle { margin-top:16px; }
  main { padding:28px 16px 64px; }
  .page-head { flex-direction:column; gap:6px; } .page-head p { margin:0; text-align:left; }
  .hero { padding:28px 22px; }
  details.file > summary { grid-template-columns:18px 1fr auto; } .fdesc, .fsize { display:none; }
  .preview { padding-left:6px; }
}
"""

JS = """
const pages = [...document.querySelectorAll('.page')];
const topics = [...document.querySelectorAll('.nav-topic')];
function pageOf(id) {
  const el = id && document.getElementById(id);
  if (!el) return null;
  return el.classList.contains('page') ? el : el.closest('.page');
}
function spy() {
  const page = pages.find(p => !p.hidden);
  if (!page) return;
  const sections = [...page.querySelectorAll('[data-spy]')];
  let current = sections[0];
  sections.forEach(s => { if (s.getBoundingClientRect().top <= 160) current = s; });
  if (window.innerHeight + window.scrollY >= document.documentElement.scrollHeight - 4) current = sections[sections.length - 1];
  document.querySelectorAll('.subnav a').forEach(a => a.classList.toggle('active', !!current && a.dataset.target === current.id));
}
function show(id, scroll) {
  const page = pageOf(id) || pages[0];
  pages.forEach(p => { p.hidden = p !== page; });
  topics.forEach(t => t.classList.toggle('open', t.dataset.page === page.id));
  const target = document.getElementById(id);
  if (window.renderMermaidIn) window.renderMermaidIn(page);
  if (scroll && target && target !== page) target.scrollIntoView({ block: 'start' });
  else window.scrollTo(0, 0);
  spy();
}
document.addEventListener('click', event => {
  const link = event.target.closest('a[href^="#"]');
  if (!link) return;
  const id = link.getAttribute('href').slice(1);
  if (!document.getElementById(id)) return;
  event.preventDefault();
  history.replaceState(null, '', '#' + id);
  show(id, true);
});
window.addEventListener('scroll', spy, { passive: true });
window.addEventListener('hashchange', () => show(location.hash.slice(1), true));
show(location.hash.slice(1), !!location.hash);
window.addEventListener('load', () => requestAnimationFrame(() => show(location.hash.slice(1), !!location.hash)));
document.querySelectorAll('details.file').forEach(d => d.addEventListener('toggle', () => {
  const frame = d.querySelector('iframe[data-src]');
  if (d.open && frame && !frame.src) frame.src = frame.dataset.src;
}));
const filter = document.getElementById('file-filter');
if (filter) filter.addEventListener('input', () => {
  const q = filter.value.toLowerCase();
  document.querySelectorAll('details.file').forEach(d => { d.style.display = d.textContent.toLowerCase().includes(q) ? '' : 'none'; });
});
const root = document.documentElement;
try { const saved = localStorage.getItem('theme'); if (saved) root.dataset.theme = saved; } catch (e) {}
document.getElementById('theme').addEventListener('click', () => {
  const dark = root.dataset.theme ? root.dataset.theme === 'dark' : matchMedia('(prefers-color-scheme: dark)').matches;
  root.dataset.theme = dark ? 'light' : 'dark';
  try { localStorage.setItem('theme', root.dataset.theme); } catch (e) {}
  location.reload();
});
"""

MERMAID_JS = """
const root = document.documentElement;
const dark = root.dataset.theme ? root.dataset.theme === 'dark' : matchMedia('(prefers-color-scheme: dark)').matches;
const light = { primaryColor:'#f3f3f3', primaryTextColor:'#000', primaryBorderColor:'#000', lineColor:'#3c3c3c',
  secondaryColor:'#ffae35', tertiaryColor:'#ffffff', clusterBkg:'#fafafa', clusterBorder:'#cecece',
  edgeLabelBackground:'#ffffff', actorBkg:'#f3f3f3', actorBorder:'#000', noteBkgColor:'#fff4e0', noteBorderColor:'#ffae35' };
const darkVars = { primaryColor:'#1f1f1f', primaryTextColor:'#fff', primaryBorderColor:'#bbbbbb', lineColor:'#bbbbbb',
  secondaryColor:'#ffae35', tertiaryColor:'#171717', clusterBkg:'#141414', clusterBorder:'#3c3c3c',
  edgeLabelBackground:'#0e0e0e', actorBkg:'#1f1f1f', actorBorder:'#bbbbbb', actorTextColor:'#fff', signalColor:'#ddd',
  signalTextColor:'#fff', noteBkgColor:'#2a2111', noteBorderColor:'#ffae35', noteTextColor:'#fff', labelTextColor:'#fff' };
mermaid.initialize({ startOnLoad:false, theme:'base', securityLevel:'loose',
  themeVariables:{ fontFamily:'Geist, Segoe UI, sans-serif', fontSize:'14px', ...(dark ? darkVars : light) },
  flowchart:{ curve:'basis', htmlLabels:true } });
window.renderMermaidIn = page => {
  const nodes = [...(page || document).querySelectorAll('pre.mermaid:not([data-processed])')];
  if (nodes.length) mermaid.run({ nodes });
};
window.renderMermaidIn(document.querySelector('.page:not([hidden])'));
"""

LIST_ITEM = re.compile(r"^(\s*)([-*]|\d+\.)\s")


def href(relative: str) -> str:
    return "../" + quote(relative)


def _normalize_lists(text: str) -> str:
    lines, previous = [], ""
    for line in text.splitlines():
        item = LIST_ITEM.match(line)
        if item and previous.strip() and not LIST_ITEM.match(previous) and not previous.lstrip().startswith("|"):
            lines.append("")
        if item and 0 < len(item.group(1)) < 4:
            line = "    " + line.lstrip()
        lines.append(line)
        previous = line
    return "\n".join(lines)


def render_markdown(text: str) -> str:
    rendered = markdown.markdown(_normalize_lists(text), extensions=["tables", "fenced_code"])
    rendered = MERMAID_BLOCK.sub(r'<pre class="mermaid">\1</pre>', rendered)
    for section in SECTIONS:
        if section.source:
            rendered = rendered.replace(f'href="{section.source}"', f'href="#{section.anchor}"')
    return rendered


def _size(path: Path) -> str:
    size = path.stat().st_size
    if size < 1024:
        return f"{size} B"
    return f"{size / 1024:.0f} KB" if size < 1_048_576 else f"{size / 1_048_576:.1f} MB"


def _text_preview(path: Path) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")
    if len(text) > TEXT_PREVIEW_LIMIT:
        text = text[:TEXT_PREVIEW_LIMIT] + "\n\n[prévia cortada; baixe o arquivo para ver tudo]"
    return f"<pre><code>{html.escape(text)}</code></pre>"


def _csv_preview(path: Path) -> str:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))[: CSV_PREVIEW_ROWS + 1]
    head = "".join(f"<th>{html.escape(c)}</th>" for c in rows[0])
    body = "".join("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in row) + "</tr>" for row in rows[1:])
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def preview(relative: str) -> str:
    path = REPO_ROOT / relative
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return (f'<p><a href="{href(relative)}" target="_blank">Abrir o PDF em outra aba</a></p>'
                f'<iframe class="frame" data-src="{href(relative)}" title="{html.escape(path.name)}"></iframe>')
    if suffix == ".png":
        return f'<img src="{href(relative)}" alt="{html.escape(path.name)}" loading="lazy">'
    if suffix == ".md":
        return render_markdown(path.read_text(encoding="utf-8"))
    if suffix == ".csv":
        return _csv_preview(path)
    if suffix in {".docx", ".xlsx"}:
        return '<p class="note">Formato sem prévia no navegador. Baixe o arquivo (a carta tem versão em PDF).</p>'
    return _text_preview(path)


def file_entry(relative: str, description: str) -> str:
    path = REPO_ROOT / relative
    if not path.exists():
        return ""
    return (
        f'<details class="file"><summary><span class="fname">{html.escape(relative)}</span>'
        f'<span class="fdesc">{html.escape(description)}</span><span class="fsize">{_size(path)}</span>'
        f'<a class="dl" href="{href(relative)}" download onclick="event.stopPropagation()">Baixar</a></summary>'
        f'<div class="preview">{preview(relative)}</div></details>'
    )


def file_sections() -> str:
    parts = ['<p>Clique no nome para ver a prévia; "Baixar" abre ou salva o arquivo. Os links funcionam com a '
             'pasta do projeto aberta no computador.</p>',
             '<div class="file-tools"><input id="file-filter" type="search" placeholder="Filtrar arquivos…"></div>']
    for anchor, title, files in FILE_GROUPS:
        entries = "".join(file_entry(relative, description) for relative, description in files)
        parts.append(f'<section class="doc files" id="{anchor}" data-spy><h2>{html.escape(title)}</h2>{entries}</section>')
    return "".join(parts)


def v1_letter_text() -> str:
    document = Document(str(OUTPUT / "output_letter.docx"))
    return "".join(f"<p>{html.escape(p.text)}</p>" for p in document.paragraphs if p.text.strip())


def letter_page_images(pdf_relative: str) -> list[str]:
    ASSETS.mkdir(exist_ok=True)
    document = pdfium.PdfDocument(str(REPO_ROOT / pdf_relative))
    names = []
    for index in range(len(document)):
        name = f"carta_pagina_{index + 1}.png"
        document[index].render(scale=1.6).to_pil().save(ASSETS / name)
        names.append(name)
    return names


def letter_section(section: Section) -> str:
    brief = (OUTPUT / "brief_assessor_albert_2025-05-07.md").read_text(encoding="utf-8")
    pdf = "Output/carta_albert_2025-05-07.pdf"
    pages = "".join(
        f'<img class="page" src="assets/{name}" alt="Carta v2, página {i}" loading="lazy">'
        for i, name in enumerate(letter_page_images(pdf), start=1)
    )
    return (
        f'<section class="doc" id="{section.anchor}" data-spy><h1>A carta e o brief</h1>'
        "<p>À esquerda, a carta que a v2 gerou para o Albert; à direita, a carta da v1, para comparar. "
        "Abaixo, o brief que o assessor recebe junto com a carta.</p>"
        f'<div class="two-col"><div><div class="col-title"><span>v2 · carta gerada</span>'
        f'<a href="{href(pdf)}" target="_blank">Abrir PDF</a></div><div class="pages">{pages}</div></div>'
        f'<div><div class="col-title"><span>v1 · carta original</span></div>'
        f'<div class="v1-letter">{v1_letter_text()}</div></div></div>'
        f'<h2>Brief do assessor</h2><div class="brief">{render_markdown(brief)}</div></section>'
    )


def _alert_count(brief: str) -> int:
    return len(re.findall(r"^\| (alta|média|info) \|", brief, flags=re.MULTILINE))


def kpi_cards() -> list[tuple[str, str]]:
    facts = json.loads((OUTPUT / "facts_albert_2025-05-07.json").read_text(encoding="utf-8"))["facts"]
    runs = json.loads((OUTPUT / "run_log_albert_2025-05-07.json").read_text(encoding="utf-8"))
    brief = (OUTPUT / "brief_assessor_albert_2025-05-07.md").read_text(encoding="utf-8")
    month, bench = facts["month"], facts["month"].get("benchmarks", {})
    idle = next(a["current"] for a in facts["allocation"] if a["bucket"].startswith("Caixa"))
    tests = sum(len(re.findall(r"^def test_", p.read_text(encoding="utf-8"), flags=re.MULTILINE))
                for p in (REPO_ROOT / "src/tests").glob("test_*.py"))
    return [
        ("Retorno no período", month["return"]),
        ("Resultado no período", month["result"]),
        ("CDI | Ibovespa", f"{bench.get('cdi', '–')} | {bench.get('ibovespa', '–')}"),
        ("Saldo + CDB vencido (% do patrimônio)", idle),
        ("Alertas para o assessor", str(_alert_count(brief))),
        ("Custo da execução", f"US$ {sum(r['cost_usd'] for r in runs):.2f}".replace(".", ",")),
        ("Grafos no Rivet", "6"),
        ("Testes automatizados", str(tests)),
    ]


def hero() -> str:
    cards = "".join(f"<div class=\"kpi\"><span>{html.escape(l)}</span><strong>{html.escape(v)}</strong></div>"
                    for l, v in kpi_cards())
    return (
        f'<div class="hero"><div class="eyebrow">{AUTHOR} · Challenge Enter · AI Deployment</div>'
        "<h2>O LLM lê e escreve, o código calcula e <em>nada chega ao cliente sem checagem.</em></h2>"
        "<p>Números da última execução para o Albert, no período de 07/04/2025 a 07/05/2025.</p>"
        f'<div class="kpis">{cards}</div>'
        '<div class="actions"><a class="btn" href="#carta">Ver a carta</a>'
        '<a class="btn secondary" href="#diagnostico">Diagnóstico da v1</a>'
        '<a class="btn secondary" href="#roteiro">Roteiro da reunião</a>'
        f'<a class="btn secondary" href="{REPO_URL}" target="_blank" rel="noopener">GitHub ↗</a></div></div>'
    )


def doc_section(section: Section) -> str:
    if section.anchor == "carta":
        return letter_section(section)
    if section.anchor == "rivet":
        return rivet_section(section.anchor)
    body = render_markdown((DOCS / section.source).read_text(encoding="utf-8"))
    return f'<section class="doc" id="{section.anchor}" data-spy>{body}</section>'


def _pager(index: int) -> str:
    links = []
    if index > 0:
        previous = TOPICS[index - 1]
        links.append(f'<a class="prev" href="#{previous.id}"><small>← Anterior</small><strong>{previous.label}</strong></a>')
    if index < len(TOPICS) - 1:
        following = TOPICS[index + 1]
        links.append(f'<a class="next" href="#{following.id}"><small>Próximo →</small><strong>{following.label}</strong></a>')
    return f'<div class="pager">{"".join(links)}</div>'


def topic_page(index: int, topic: Topic) -> str:
    head = (f'<div class="page-head"><span class="num">{index + 1:02d}</span><h1>{topic.label}</h1>'
            f"<p>{html.escape(topic.summary)}</p></div>")
    if topic.id == "arquivos":
        body = file_sections()
    else:
        body = "".join(doc_section(s) for s in SECTIONS if s.topic == topic.id)
    intro = hero() if topic.id == "apresentacao" else ""
    return f'<div class="page" id="{topic.id}">{head}{intro}{body}{_pager(index)}</div>'


def _subnav(topic: Topic) -> list[tuple[str, str]]:
    if topic.id == "arquivos":
        return [(anchor, title) for anchor, title, _ in FILE_GROUPS]
    return [(s.anchor, s.nav) for s in SECTIONS if s.topic == topic.id]


def sidebar() -> str:
    parts = ['<nav class="sidebar"><div class="brand"><div class="mark">ENTER<b></b></div>'
             "<small>Challenge · Carta mensal XP com IA</small>"
             f'<div class="author"><strong>{AUTHOR}</strong>'
             f'<a href="{REPO_URL}" target="_blank" rel="noopener">GitHub · enter-challenge ↗</a></div></div>']
    for index, topic in enumerate(TOPICS):
        items = "".join(f'<a href="#{anchor}" data-target="{anchor}">{html.escape(title)}</a>' for anchor, title in _subnav(topic))
        parts.append(
            f'<div class="nav-topic" data-page="{topic.id}"><a class="topic-link" href="#{topic.id}">'
            f'<span class="num">{index + 1:02d}</span>{topic.label}</a><div class="subnav">{items}</div></div>'
        )
    parts.append('<button class="theme-toggle" id="theme" type="button">Tema claro / escuro</button></nav>')
    return "".join(parts)


def mermaid_script() -> str:
    return (
        f"import mermaid from '{MERMAID_URL}';" + MERMAID_JS
    )


def build() -> Path:
    topbar = (f'<div class="topbar">{AUTHOR} · Carta mensal XP com IA: da v1 à v2, com números calculados e checados'
              f'<a href="{REPO_URL}" target="_blank" rel="noopener">GitHub</a><a href="#carta">Ver a carta</a></div>')
    pages = "".join(topic_page(i, topic) for i, topic in enumerate(TOPICS))
    page = (
        '<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        f'<title>Carta mensal XP · Documentação</title><link rel="preconnect" href="https://fonts.googleapis.com">'
        f'<link href="{FONTS}" rel="stylesheet"><style>{CSS}{RIVET_CSS}</style></head>'
        f'<body>{topbar}<div class="layout">{sidebar()}<main>{pages}</main></div>'
        f'<script>{JS}{RIVET_JS}</script><script type="module">{mermaid_script()}</script></body></html>'
    )
    target = DOCS / "index.html"
    target.write_text(page, encoding="utf-8")
    return target


if __name__ == "__main__":
    print(build())
