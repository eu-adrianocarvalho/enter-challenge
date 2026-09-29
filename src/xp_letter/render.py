"""Monta a carta em DOCX com python-docx, 100% por código: cabeçalho, destinatário, assunto, texto,
indicadores, gráfico, tabela de alocação, tabela de sugestões, observações, assinatura e disclaimer.
Converte para PDF abrindo uma instância separada do Word (ou o LibreOffice, se existir) e conta as
páginas para garantir o limite de duas."""
from __future__ import annotations

import shutil
import subprocess
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pypdf import PdfReader

from xp_letter.facts import Recommendation

BLACK_HEX = "000000"
# Amarelo aproximado da XP: os sites da XP bloqueiam acesso automatizado, então conferir no brand book.
XP_YELLOW_HEX = "FFC709"
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x6B, 0x6B, 0x6B)
LIGHT_GRAY = RGBColor(0xBF, 0xBF, 0xBF)
XP_YELLOW = RGBColor.from_string(XP_YELLOW_HEX)
BODY_FONT = "Calibri"
WORD_PDF_FORMAT = 17


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _remove_borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    table._tbl.tblPr.append(borders)


def _cell_bottom_border(cell, color: str, size: int) -> None:
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", str(size)), ("space", "0"), ("color", color)):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    cell._tc.get_or_add_tcPr().append(borders)


def _border(paragraph, side: str, color: str, size: int, space: int) -> None:
    border = OxmlElement("w:pBdr")
    edge = OxmlElement(f"w:{side}")
    for key, value in (("val", "single"), ("sz", str(size)), ("space", str(space)), ("color", color)):
        edge.set(qn(f"w:{key}"), value)
    border.append(edge)
    paragraph._p.get_or_add_pPr().append(border)


def _run(paragraph, text: str, size: float = 10, bold: bool = False, color: RGBColor | None = None):
    run = paragraph.add_run(text)
    run.font.size, run.font.bold = Pt(size), bold
    if color is not None:
        run.font.color.rgb = color
    return run


def _setup(document) -> None:
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.top_margin, section.bottom_margin = Cm(1.4), Cm(1.3)
    section.left_margin = section.right_margin = Cm(2.0)
    section.header_distance = section.footer_distance = Cm(0.7)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = BODY_FONT, Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12


def _brand_band(document, logo: Path) -> None:
    if not logo.exists():
        raise FileNotFoundError(f"Logo não encontrado: {logo}. Ajuste brand.logo em config/settings.yaml.")
    header = document.sections[0].header
    table = header.add_table(rows=1, cols=2, width=Cm(17))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _remove_borders(table)
    logo_cell, title_cell = table.rows[0].cells
    for cell, width in ((logo_cell, Cm(6)), (title_cell, Cm(11))):
        _shade(cell, BLACK_HEX)
        _cell_bottom_border(cell, XP_YELLOW_HEX, 24)
        cell.width, cell.vertical_alignment = width, WD_ALIGN_VERTICAL.CENTER
    table.rows[0].height, table.rows[0].height_rule = Cm(1.4), WD_ROW_HEIGHT_RULE.EXACTLY
    for cell in (logo_cell, title_cell):
        spacing = cell.paragraphs[0].paragraph_format
        spacing.space_before, spacing.space_after, spacing.line_spacing = Pt(0), Pt(0), 1.0
    logo_cell.paragraphs[0].add_run().add_picture(str(logo), height=Cm(0.9))
    title = title_cell.paragraphs[0]
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(title, "Relatório mensal de investimentos", 9.5, bold=True, color=WHITE)
    _run(title, "\nAssessoria de Investimentos", 8, color=LIGHT_GRAY)
    spacer = header.paragraphs[0]
    spacer._p.addprevious(table._tbl)
    spacer.paragraph_format.space_after = Pt(0)


def _header(document, facts: dict[str, Any], logo: Path) -> None:
    _brand_band(document, logo)
    date_line = document.add_paragraph()
    date_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(date_line, facts["letter_date"], 9, color=GRAY)
    recipient = document.add_paragraph()
    recipient.paragraph_format.space_after = Pt(2)
    _run(recipient, facts["client"]["full_name"], 10, bold=True)
    _run(recipient, f"\nConta {facts['client']['account']} · Perfil {facts['profile']['label']}", 9, color=GRAY)


def _footer(document, disclaimer: str) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _run(paragraph, disclaimer, 6.5, color=GRAY)


def _body(document, text: str, new_page: bool = False, keep_with_next: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.page_break_before = new_page
    paragraph.paragraph_format.keep_with_next = keep_with_next
    _run(paragraph, text.strip(), 10.5)


def _kpis(document, facts: dict[str, Any]) -> None:
    month = facts["month"]
    bench = month.get("benchmarks", {})
    boxes = [
        ("Patrimônio total", facts["wealth"]["net_worth"]),
        ("Retorno no período*", month["return"]),
        ("Resultado no período*", month["result"]),
        ("CDI  |  Ibovespa", f"{bench.get('cdi', '–')}  |  {bench.get('ibovespa', '–')}"),
    ]
    table = document.add_table(rows=1, cols=len(boxes))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _remove_borders(table)
    for cell, (label, value) in zip(table.rows[0].cells, boxes):
        _shade(cell, BLACK_HEX)
        top = cell.paragraphs[0]
        top.paragraph_format.space_after = Pt(0)
        _run(top, label, 7.5, color=LIGHT_GRAY)
        _run(cell.add_paragraph(), value, 12, bold=True, color=XP_YELLOW)
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(2)
    _run(note, f"*Período de {facts['period']['start']} a {facts['period']['end']}, considerando {month['covers']}.", 7, color=GRAY)


def _chart(document, chart: Path) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run().add_picture(str(chart), width=Cm(15.0))


def _allocation_table(document, facts: dict[str, Any]) -> None:
    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(3)
    _run(title, f"Sua alocação hoje e a faixa indicada para o perfil {facts['profile']['label']}", 9.5, bold=True, color=BLACK)
    rows = facts["allocation"]
    table = document.add_table(rows=1 + len(rows), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, ("Classe", "Hoje (% do patrimônio)", "Faixa indicada")):
        _shade(cell, BLACK_HEX)
        _run(cell.paragraphs[0], header, 8.5, bold=True, color=WHITE)
    for row, allocation in zip(table.rows[1:], rows):
        for index, value in enumerate((allocation["bucket"], allocation["current"], allocation["target_band"])):
            paragraph = row.cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            outside = index == 1 and allocation["status"] != "dentro"
            _run(paragraph, value, 8.5, bold=outside)


def _recommendation_table(document, recommendations: list[Recommendation]) -> None:
    rows = [(r.title if index == 0 else "", line) for r in recommendations for index, line in enumerate(r.lines())]
    table = document.add_table(rows=1 + len(rows), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = (Cm(6.2), Cm(2.4), Cm(5.2), Cm(3.2))
    for cell, title, width in zip(table.rows[0].cells, ("Sugestão", "Movimento", "Ativo", "Valor"), widths):
        _shade(cell, BLACK_HEX)
        cell.width = width
        _run(cell.paragraphs[0], title, 8.5, bold=True, color=WHITE)
    for row, (title, line) in zip(table.rows[1:], rows):
        values = (title, line["action"], line["asset"], line["amount"])
        for index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = width
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            if index == 3:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _run(paragraph, value, 8.5, bold=index == 0)


def _signature(document, facts: dict[str, Any]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    _run(paragraph, "Atenciosamente,\n")
    _run(paragraph, facts["advisor"]["name"], 10, bold=True)
    _run(paragraph, f"\nAssessor de Investimentos XP · Código {facts['advisor']['code']}", 9, color=GRAY)


def disclaimer_text(facts: dict[str, Any]) -> str:
    return (
        "Este material tem caráter informativo e não constitui oferta ou solicitação de compra ou venda de ativos. "
        "As sugestões dependem da confirmação de adequação ao seu perfil (suitability) e da sua autorização prévia. "
        "Rentabilidade passada não é garantia de rentabilidade futura; valores brutos de impostos. Retorno do período "
        "calculado com preços de fechamento e cotas diárias publicadas pela CVM. Fontes: extrato XP de "
        f"{facts['period']['end']}, CVM, Banco Central do Brasil, Yahoo Finance e {facts['macro']['source']}."
    )


def _table_notes(document, notes: list[str]) -> None:
    for note in notes:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(2)
        _run(paragraph, f"Observação: {note}", 8, color=GRAY)


def render_docx(letter: dict[str, str], facts: dict[str, Any], recommendations: list[Recommendation],
                notes: list[str], chart: Path, logo: Path, target: Path) -> Path:
    document = Document()
    _setup(document)
    _header(document, facts, logo)
    _footer(document, disclaimer_text(facts))
    subject = document.add_paragraph()
    _border(subject, "left", XP_YELLOW_HEX, 24, 6)
    _run(subject, letter["subject"].strip(), 11.5, bold=True, color=BLACK)
    _body(document, letter["greeting"])
    _body(document, letter["performance"])
    _kpis(document, facts)
    _chart(document, chart)
    _allocation_table(document, facts)
    _body(document, letter["outlook"], new_page=True)
    _body(document, letter["recommendations"])
    if recommendations:
        _recommendation_table(document, recommendations)
        _table_notes(document, notes)
        document.add_paragraph().paragraph_format.space_after = Pt(0)
    _body(document, letter["closing"], keep_with_next=True)
    _signature(document, facts)
    document.save(target)
    return target


def _word_to_pdf(docx: Path, pdf: Path) -> bool:
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return False
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible, word.DisplayAlerts = False, 0
    try:
        document = word.Documents.Open(str(docx.resolve()), ReadOnly=True)
        document.SaveAs(str(pdf.resolve()), FileFormat=WORD_PDF_FORMAT)
        document.Close(False)
        return True
    finally:
        word.Quit()
        pythoncom.CoUninitialize()


def _libreoffice_to_pdf(docx: Path, pdf: Path) -> bool:
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office is None:
        return False
    subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", str(pdf.parent), str(docx)], check=True)
    return pdf.exists()


def docx_to_pdf(docx: Path) -> Path | None:
    pdf = docx.with_suffix(".pdf")
    return pdf if _word_to_pdf(docx, pdf) or _libreoffice_to_pdf(docx, pdf) else None


def page_count(pdf: Path) -> int:
    return len(PdfReader(str(pdf)).pages)
